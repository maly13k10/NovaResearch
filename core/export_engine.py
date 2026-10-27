from __future__ import annotations

import io
import json
import os
import platform
import sys
import zipfile
from datetime import datetime
from typing import Any

import pandas as pd
from docx import Document


def environment_snapshot() -> dict[str, Any]:
    packages = {}
    for name in ["streamlit", "pandas", "numpy", "matplotlib", "sklearn", "pypdf", "docx", "scipy", "networkx", "shap", "diffprivlib", "openai"]:
        try:
            module = __import__(name)
            packages[name] = getattr(module, "__version__", "installed")
        except Exception:
            packages[name] = "not installed"
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "python": sys.version,
        "platform": platform.platform(),
        "packages": packages,
    }


def df_records(df) -> Any:
    if df is None:
        return None
    try:
        return df.replace({float("nan"): None}).to_dict("records")
    except Exception:
        return None


def serialize_state(state: dict[str, Any]) -> dict[str, Any]:
    project = state.get("project")
    return {
        "project": project,
        "active_idea_idx": state.get("active_idea_idx", 0),
        "novelty": state.get("novelty"),
        "experiment_plan": state.get("experiment_plan"),
        "selected_journal": state.get("selected_journal"),
        "research_story": state.get("research_story"),
        "publication_ready": bool(state.get("publication_ready")),
        "activity_log": state.get("activity_log", []),
        "active_target": state.get("active_target"),
        "repo_url": state.get("repo_url", ""),
        "simulator_version": state.get("simulator_version", ""),
        "availability_statement": state.get("availability_statement", ""),
        "statistical_validation": state.get("statistical_validation"),
        "dp_note": state.get("dp_note", ""),
        "environment_snapshot": state.get("environment_snapshot") or environment_snapshot(),
    }


def build_session_archive(state: dict[str, Any]) -> bytes:
    buff = io.BytesIO()
    with zipfile.ZipFile(buff, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("state.json", json.dumps(serialize_state(state), ensure_ascii=False, indent=2, default=str))
        tables = {
            "data/active_df.csv": state.get("active_df"),
            "results/ml_results.csv": state.get("ml_results"),
            "results/dp_results.csv": state.get("dp_results"),
            "results/importance_results.csv": state.get("importance_results"),
            "results/ablation_results.csv": state.get("ablation_results"),
            "results/drift_results.csv": state.get("drift_results"),
            "results/shap_global.csv": state.get("shap_global"),
            "results/shap_local.csv": state.get("shap_local"),
        }
        rob = state.get("robustness_results") or {}
        tables["results/robustness_noise.csv"] = rob.get("noise")
        tables["results/robustness_missing.csv"] = rob.get("missing")
        for path, df in tables.items():
            if df is not None:
                z.writestr(path, df.to_csv(index=False))
    return buff.getvalue()


def parse_session_archive(payload: bytes) -> dict[str, Any]:
    restored: dict[str, Any] = {}
    with zipfile.ZipFile(io.BytesIO(payload), "r") as z:
        state = json.loads(z.read("state.json").decode("utf-8"))
        restored.update(state)
        mapping = {
            "data/active_df.csv": "active_df", "results/ml_results.csv": "ml_results",
            "results/dp_results.csv": "dp_results", "results/importance_results.csv": "importance_results",
            "results/ablation_results.csv": "ablation_results", "results/drift_results.csv": "drift_results",
            "results/shap_global.csv": "shap_global", "results/shap_local.csv": "shap_local",
        }
        for path, key in mapping.items():
            if path in z.namelist():
                restored[key] = pd.read_csv(io.BytesIO(z.read(path)))
        noise = pd.read_csv(io.BytesIO(z.read("results/robustness_noise.csv"))) if "results/robustness_noise.csv" in z.namelist() else None
        missing = pd.read_csv(io.BytesIO(z.read("results/robustness_missing.csv"))) if "results/robustness_missing.csv" in z.namelist() else None
        restored["robustness_results"] = {"noise": noise, "missing": missing} if noise is not None or missing is not None else None
    return restored


def generate_paper_sections(context: dict[str, Any], state: dict[str, Any]) -> dict[str, str]:
    project = state.get("project") or {}
    idea = context.get("active_idea") or {}
    novelty = state.get("novelty") or {}
    ml = state.get("ml_results")
    dp = state.get("dp_results")
    rob = state.get("robustness_results") or {}
    stats = state.get("statistical_validation") or {}
    best_model = "not yet evaluated"
    best_f1 = "—"
    if ml is not None and len(ml):
        valid = ml.dropna(subset=["F1"])
        if len(valid):
            best_model = str(valid.iloc[0]["Model"]); best_f1 = f"{valid.iloc[0]['F1']:.3f}"
    dp_summary = "Differential Privacy has not yet been evaluated."
    if dp is not None and len(dp):
        best = dp.dropna(subset=["F1"]).sort_values("F1", ascending=False).iloc[0]
        dp_summary = f"The highest observed private-model F1 was {best['F1']:.3f} at ε={best['Epsilon']}. {state.get('dp_note','')}"
    noise_summary = "Robustness has not yet been evaluated."
    if rob.get("noise") is not None and len(rob["noise"]):
        noise_summary = f"Noise robustness ranged from F1={rob['noise']['F1'].max():.3f} at baseline to F1={rob['noise']['F1'].iloc[-1]:.3f} at the highest tested noise level."
    references = []
    for work in novelty.get("similar", [])[:10]:
        references.append(f"{work.get('Title','Untitled')} ({work.get('Year','—')}). DOI: {work.get('DOI','—')}")
    return {
        "Abstract": f"This study investigates {idea.get('title','the active research problem')} in the domain of {context.get('domain')}. The workflow combines reproducible data preparation, baseline machine learning, Differential Privacy, robustness, explainability and statistical validation. The current best benchmark model is {best_model} with weighted F1={best_f1}. The novelty-screening indicator is {novelty.get('Overall novelty','not analyzed')}/100. Final claims require verification against the completed experiments and literature.",
        "Introduction & Research Gap": f"The project '{context.get('project_title')}' addresses the research question: {idea.get('question','—')}. Automated readiness analysis identified these structural gaps: {', '.join(context.get('missing') or []) or 'none in the initial screening'}. The intended contribution is: {idea.get('novelty_hook','—')}",
        "Related Work": "The novelty module retrieves literature candidates through Crossref and compares title-level similarity. The closest candidates should be reviewed manually to establish the final research gap.\n\n" + "\n".join(references),
        "Dataset and Input Variables": f"Candidate datasets include: {', '.join(context.get('datasets') or [])}. Active inputs are {', '.join(idea.get('inputs',[]))}; target: {idea.get('target','—')}. Data provenance, license and preprocessing should be documented in the final manuscript.",
        "Simulator / Experimental Environment": f"The active simulator/environment is {idea.get('simulator','—')}. Recorded simulator version: {state.get('simulator_version') or 'pending'}. Environment metadata is exported with the publication package.",
        "ML and Differential Privacy Methods": f"The benchmark compares Logistic Regression, Random Forest, Gradient Boosting and SVM-RBF. {dp_summary}",
        "Experimental Protocol": json.dumps(state.get("experiment_plan") or {}, ensure_ascii=False, indent=2),
        "Results": f"The current best non-private model is {best_model} with weighted F1={best_f1}. Tables exported by NovaResearch AI contain the full benchmark and privacy results.",
        "Explainability & Robustness": f"{noise_summary} Explainability uses permutation importance plus global/local SHAP when supported. Automatic ablation and sensor-drift analyses quantify feature dependence and degradation.",
        "Privacy–Utility Analysis": dp_summary,
        "Statistical Validation": f"Exact paired McNemar-style validation: {stats if stats else 'not yet run'}. A p-value below 0.05 is treated as evidence of different error patterns between the compared models.",
        "Discussion & Limitations": "Automated scores are decision-support indicators rather than proof of novelty or publication quality. Limitations include dataset representativeness, simulator-to-real transfer, model assumptions, privacy-mechanism assumptions and external-validity constraints.",
        "Conclusions": f"The workflow provides an integrated path from project understanding to reproducible evidence. The next recommended action is {context.get('next_action_title')}: {context.get('next_action_text')}",
        "Data Availability": state.get("availability_statement") or "Data availability statement pending. Document the source, license and generated-data procedure before submission.",
        "Code Availability": f"Repository: {state.get('repo_url') or 'pending'}. The final package exports environment and reproducibility metadata.",
    }


def build_paper_docx(sections: dict[str, str], title: str, subtitle: str = "NovaResearch AI publication draft") -> bytes:
    doc = Document()
    doc.add_heading(title or "NovaResearch AI Research Draft", 0)
    doc.add_paragraph(subtitle)
    for heading, body in sections.items():
        doc.add_heading(heading, level=1)
        doc.add_paragraph(str(body))
    doc.add_heading("Research Integrity Note", level=1)
    doc.add_paragraph("NovaResearch AI provides research decision support. Verify all citations, novelty claims, venue requirements, statistical assumptions and privacy guarantees before submission.")
    buff = io.BytesIO(); doc.save(buff); return buff.getvalue()


def build_publication_package(state: dict[str, Any], context: dict[str, Any], sections: dict[str, str], passport: pd.DataFrame, plots: dict[str, bytes] | None = None) -> bytes:
    buff = io.BytesIO()
    env = state.get("environment_snapshot") or environment_snapshot()
    with zipfile.ZipFile(buff, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("01_project_state.json", json.dumps(serialize_state(state), ensure_ascii=False, indent=2, default=str))
        z.writestr("02_research_story.md", state.get("research_story") or "")
        z.writestr("03_experiment_plan.json", json.dumps(state.get("experiment_plan") or {}, ensure_ascii=False, indent=2, default=str))
        z.writestr("04_journal_match.json", json.dumps(state.get("selected_journal") or {}, ensure_ascii=False, indent=2, default=str))
        z.writestr("05_reproducibility_passport.csv", passport.to_csv(index=False))
        z.writestr("06_environment.json", json.dumps(env, ensure_ascii=False, indent=2, default=str))
        z.writestr("07_paper_sections.json", json.dumps(sections, ensure_ascii=False, indent=2))
        z.writestr("08_publication_draft.docx", build_paper_docx(sections, context.get("project_title", "NovaResearch AI Research Draft")))
        z.writestr("09_requirements_snapshot.txt", "\n".join(f"{k}=={v}" for k, v in env.get("packages", {}).items()))
        tables = {
            "results/ml_benchmark.csv": state.get("ml_results"), "results/differential_privacy.csv": state.get("dp_results"),
            "results/permutation_importance.csv": state.get("importance_results"), "results/ablation.csv": state.get("ablation_results"),
            "results/drift.csv": state.get("drift_results"), "results/shap_global.csv": state.get("shap_global"), "results/shap_local.csv": state.get("shap_local"),
        }
        rob = state.get("robustness_results") or {}; tables["results/robustness_noise.csv"] = rob.get("noise"); tables["results/robustness_missing.csv"] = rob.get("missing")
        for path, df in tables.items():
            if df is not None:
                z.writestr(path, df.to_csv(index=False))
        if plots:
            for name, payload in plots.items():
                if payload:
                    z.writestr(f"plots/{name}.png", payload)
        z.writestr("README.txt", "NovaResearch AI final publication package. Verify citations, venue requirements, statistical assumptions and formal privacy claims before submission.\n")
    return buff.getvalue()
